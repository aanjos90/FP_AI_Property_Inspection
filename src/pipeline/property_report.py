"""Makes the report for a property one section per room"""

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Cm, Pt

from src.pipeline.schema import SECTION_NAMES, SECTIONS, room_display_name

PROPERTIES_DIR = Path("properties")

FONT_NAME = "Arial"
BODY_SIZE = Pt(12)
TITLE_SIZE = Pt(14)
TITLE = "MOVE-IN INSPECTION REPORT"

# A4 page
PAGE_WIDTH = Cm(21)
PAGE_HEIGHT = Cm(29.7)
MARGIN_TOP_BOTTOM = Cm(2.54)
MARGIN_LEFT_RIGHT = Cm(1.91)


def _set_page_layout(document):
    page = document.sections[0]
    page.page_width = PAGE_WIDTH
    page.page_height = PAGE_HEIGHT
    page.top_margin = MARGIN_TOP_BOTTOM
    page.bottom_margin = MARGIN_TOP_BOTTOM
    page.left_margin = MARGIN_LEFT_RIGHT
    page.right_margin = MARGIN_LEFT_RIGHT


def _add_run(paragraph, text: str, bold: bool = False, size=BODY_SIZE):
    run = paragraph.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = size
    run.bold = bold
    return run


def generate_property_report(property_id: str, drafts: list):
    """Build the property report
        property_id  the ID of the property
        drafts  one draft per room in the order the rooms should appear
    """
    document = Document()
    _set_page_layout(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(title, TITLE, bold=True, size=TITLE_SIZE)

    reference = document.add_paragraph()
    _add_run(reference, "PROPERTY ID: ", bold=True)
    _add_run(reference, property_id)
    document.add_paragraph()

    room_number = 1
    for draft in drafts:
        # Room name
        room_name = room_display_name(draft["room"]).upper()
        heading = document.add_paragraph()
        heading_text = _add_run(heading, f"{room_number}. {room_name}:", bold=True)
        heading_text.underline = True
        room_number = room_number + 1

        # One line for each section of the room that was observed
        for section in SECTIONS:
            if section not in draft["sections"]:
                continue  # not observed in this room, skip it
            entry = draft["sections"][section]

            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            label = SECTION_NAMES[section].capitalize()
            label_run = _add_run(paragraph, f"{label}: ", bold=True)
            text_run = _add_run(paragraph, entry["text"])

            # Uncertain line highlighted yellow
            if entry["flagged_for_review"]:
                label_run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                text_run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    return document


def get_room_drafts(property_id: str) -> list:
    outputs_dir = PROPERTIES_DIR / property_id / "outputs"
    if not outputs_dir.exists():
        return []

    drafts = []
    for draft_file in sorted(outputs_dir.glob("*/draft.json")):
        drafts.append(json.loads(draft_file.read_text(encoding="utf-8")))
    return drafts


def save_property_report(property_id: str) -> Path:
    """Make the property report from the room drafts and save it"""
    drafts = get_room_drafts(property_id)
    if len(drafts) == 0:
        raise ValueError(f"No finalised drafts found for property {property_id}")

    document = generate_property_report(property_id, drafts)
    out_file = PROPERTIES_DIR / property_id / "outputs" / f"{property_id}_report.docx"
    document.save(str(out_file))
    return out_file
